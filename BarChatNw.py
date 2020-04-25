import  openpyxl as wb
from openpyxl.chart import BarChart,Reference
myshet=wb.load_workbook('Book1.xlsx')
sheet=myshet['Sheet1']
value=Reference(sheet,
                min_row=2,
                max_row=sheet.max_row,
                min_col=1,
                max_col=1)
chart=BarChart()
chart.add_data(value)
chart.title="My BarChart"
chart.x_axis.title="Time"
chart.y_axis.title="Date"
sheet.add_chart(chart,'b2')
myshet.save('NewBook1.xlsx')
